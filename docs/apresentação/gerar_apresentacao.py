from io import BytesIO
from pathlib import Path

from PIL import Image, ImageEnhance, ImageOps
from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_CONNECTOR, MSO_SHAPE
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
from pptx.util import Inches, Pt


ROOT = Path(__file__).resolve().parents[2]
OUT = Path(__file__).resolve().parent / "Ecos-de-uma-Civilizacao.pptx"

IMAGES = {
    index: ROOT / "colecao" / f"{index:02d}-descoberta" / "imagem.png"
    for index in range(1, 13)
}
DISCARD = ROOT / "descartes" / "imagens" / "10-cozinha-salto-narrativo.png"

BG = RGBColor(20, 24, 21)
PANEL = RGBColor(31, 37, 32)
PANEL_2 = RGBColor(39, 45, 39)
STONE = RGBColor(236, 229, 214)
MUTED = RGBColor(181, 180, 164)
SAND = RGBColor(198, 178, 132)
COPPER = RGBColor(75, 151, 141)
COPPER_DARK = RGBColor(43, 99, 92)
RUST = RGBColor(173, 103, 66)
BLACK = RGBColor(9, 11, 10)

TITLE_FONT = "Georgia"
BODY_FONT = "Aptos"

prs = Presentation()
prs.slide_width = Inches(13.333333)
prs.slide_height = Inches(7.5)
blank = prs.slide_layouts[6]
_streams = []


def inch(value):
    return Inches(value)


def set_background(slide, color=BG):
    fill = slide.background.fill
    fill.solid()
    fill.fore_color.rgb = color


def add_rect(slide, x, y, w, h, color, line=None, radius=False):
    kind = MSO_SHAPE.ROUNDED_RECTANGLE if radius else MSO_SHAPE.RECTANGLE
    shape = slide.shapes.add_shape(kind, inch(x), inch(y), inch(w), inch(h))
    shape.fill.solid()
    shape.fill.fore_color.rgb = color
    shape.line.color.rgb = line if line else color
    if radius:
        shape.adjustments[0] = 0.08
    return shape


def add_line(slide, x1, y1, x2, y2, color=COPPER, width=1.5):
    line = slide.shapes.add_connector(
        MSO_CONNECTOR.STRAIGHT, inch(x1), inch(y1), inch(x2), inch(y2)
    )
    line.line.color.rgb = color
    line.line.width = Pt(width)
    return line


def add_text(
    slide,
    text,
    x,
    y,
    w,
    h,
    size=18,
    color=STONE,
    bold=False,
    font=BODY_FONT,
    align=PP_ALIGN.LEFT,
    valign=MSO_ANCHOR.TOP,
    margin=0,
    line_spacing=1.0,
):
    box = slide.shapes.add_textbox(inch(x), inch(y), inch(w), inch(h))
    frame = box.text_frame
    frame.clear()
    frame.word_wrap = True
    frame.margin_left = inch(margin)
    frame.margin_right = inch(margin)
    frame.margin_top = inch(margin)
    frame.margin_bottom = inch(margin)
    frame.vertical_anchor = valign
    paragraph = frame.paragraphs[0]
    paragraph.text = text
    paragraph.alignment = align
    paragraph.line_spacing = line_spacing
    paragraph.space_after = Pt(0)
    for run in paragraph.runs:
        run.font.name = font
        run.font.size = Pt(size)
        run.font.bold = bold
        run.font.color.rgb = color
    return box


def add_rich_text(slide, segments, x, y, w, h, size=18, align=PP_ALIGN.LEFT):
    box = slide.shapes.add_textbox(inch(x), inch(y), inch(w), inch(h))
    frame = box.text_frame
    frame.clear()
    frame.word_wrap = True
    frame.margin_left = 0
    frame.margin_right = 0
    frame.margin_top = 0
    frame.margin_bottom = 0
    paragraph = frame.paragraphs[0]
    paragraph.alignment = align
    paragraph.space_after = Pt(0)
    for text_value, color, bold, font_name in segments:
        run = paragraph.add_run()
        run.text = text_value
        run.font.name = font_name
        run.font.size = Pt(size)
        run.font.bold = bold
        run.font.color.rgb = color
    return box


def image_stream(path, width_px, height_px, darken=1.0, center=(0.5, 0.5)):
    with Image.open(path) as source:
        image = source.convert("RGB")
        image = ImageOps.fit(
            image,
            (width_px, height_px),
            method=Image.Resampling.LANCZOS,
            centering=center,
        )
        if darken != 1.0:
            image = ImageEnhance.Brightness(image).enhance(darken)
        stream = BytesIO()
        image.save(stream, format="JPEG", quality=91, optimize=True)
        stream.seek(0)
        _streams.append(stream)
        return stream


def add_image(slide, path, x, y, w, h, darken=1.0, center=(0.5, 0.5), border=None):
    pixel_w = max(300, int(w * 180))
    pixel_h = max(300, int(h * 180))
    stream = image_stream(path, pixel_w, pixel_h, darken=darken, center=center)
    picture = slide.shapes.add_picture(stream, inch(x), inch(y), inch(w), inch(h))
    if border:
        picture.line.color.rgb = border
        picture.line.width = Pt(1)
    return picture


def add_footer(slide, number):
    add_line(slide, 0.45, 7.18, 12.88, 7.18, COPPER_DARK, 0.8)
    add_text(slide, "ECOS DE UMA CIVILIZAÇÃO", 0.48, 7.20, 3.5, 0.18, 7.5, MUTED, True)
    add_text(slide, f"{number:02d}", 12.25, 7.19, 0.55, 0.2, 8, COPPER, True, align=PP_ALIGN.RIGHT)


def add_slide_title(slide, eyebrow, title, number):
    add_text(slide, eyebrow.upper(), 0.48, 0.30, 4.5, 0.22, 8.5, COPPER, True)
    add_text(slide, title, 0.48, 0.55, 12.2, 0.55, 25, STONE, True, TITLE_FONT)
    add_footer(slide, number)


def add_number_badge(slide, number, x, y):
    add_rect(slide, x, y, 0.48, 0.30, COPPER, COPPER, radius=True)
    add_text(
        slide,
        f"{number:02d}",
        x,
        y + 0.01,
        0.48,
        0.23,
        9,
        BLACK,
        True,
        align=PP_ALIGN.CENTER,
        valign=MSO_ANCHOR.MIDDLE,
    )


def add_discovery_card(slide, index, name, summary, x):
    add_rect(slide, x, 1.25, 4.05, 5.68, PANEL, COPPER_DARK, radius=True)
    add_image(slide, IMAGES[index], x + 0.10, 1.35, 3.85, 3.98, border=RGBColor(65, 77, 67))
    add_number_badge(slide, index, x + 0.22, 1.50)
    add_text(slide, name, x + 0.17, 5.48, 3.70, 0.45, 14.2, STONE, True, TITLE_FONT)
    add_text(slide, summary, x + 0.17, 5.95, 3.70, 0.79, 10.6, MUTED, False, BODY_FONT, line_spacing=0.92)


def new_slide():
    slide = prs.slides.add_slide(blank)
    set_background(slide)
    return slide


# 1 — Capa
slide = new_slide()
add_image(slide, IMAGES[1], 6.95, 0, 6.383, 7.5, darken=0.78, center=(0.50, 0.45))
add_rect(slide, 0, 0, 7.18, 7.5, BG)
add_line(slide, 0.62, 1.12, 3.20, 1.12, COPPER, 2.2)
add_text(slide, "ECOS DE UMA", 0.62, 1.55, 6.05, 0.75, 31, STONE, True, TITLE_FONT)
add_text(slide, "CIVILIZAÇÃO", 0.62, 2.25, 6.05, 0.80, 31, COPPER, True, TITLE_FONT)
add_text(
    slide,
    "Uma arqueologia fictícia construída\ndescoberta por descoberta",
    0.65,
    3.28,
    5.75,
    1.05,
    18,
    SAND,
    False,
    TITLE_FONT,
    line_spacing=1.0,
)
add_text(
    slide,
    "Criatividade Computacional · CIn/UFPE · 2026.2",
    0.65,
    5.52,
    5.8,
    0.32,
    10.5,
    MUTED,
    True,
)
add_text(
    slide,
    "Ellian Rodrigues · Fabriely Santos · Aline Marianna\nMonyque Lima · Amanda Arruda · Nícolas Veiga",
    0.65,
    5.93,
    5.95,
    0.70,
    10,
    STONE,
    False,
    line_spacing=0.95,
)


# 2 — Proposta
slide = new_slide()
add_slide_title(slide, "Projeto", "Uma história que emerge das descobertas", 2)
add_text(
    slide,
    "Um museu virtual de uma civilização desaparecida cuja história não foi escrita previamente.",
    0.55,
    1.45,
    5.5,
    1.35,
    24,
    STONE,
    True,
    TITLE_FONT,
)
add_text(
    slide,
    "Cada artefato altera o que sabemos e passa a fazer parte do contexto da próxima geração.",
    0.58,
    3.08,
    5.25,
    1.15,
    16,
    SAND,
)
add_rect(slide, 0.58, 4.68, 5.10, 1.15, PANEL, COPPER_DARK, radius=True)
add_text(slide, "Não ilustramos uma história pronta.", 0.84, 4.91, 4.6, 0.28, 13, MUTED, False)
add_text(slide, "Construímos a história durante o processo.", 0.84, 5.25, 4.6, 0.34, 15, COPPER, True)
for idx, image_index in enumerate((1, 4, 12)):
    x = 6.30 + idx * 2.25
    add_image(slide, IMAGES[image_index], x, 1.34, 2.05, 5.45, darken=0.92, border=RGBColor(68, 78, 69))
    add_number_badge(slide, image_index, x + 0.13, 1.48)


# 3 — Eixo
slide = new_slide()
add_slide_title(slide, "Eixo", "A regra que transforma imagens em coleção", 3)
add_image(slide, IMAGES[2], 9.45, 1.23, 3.42, 5.70, darken=0.82, border=RGBColor(63, 76, 67))
axis_cards = [
    ("RESTRIÇÃO", "Cada descoberta precisa ser coerente com o conhecimento acumulado."),
    ("O QUE SE REPETE", "Toda imagem acrescenta informação relevante sobre a mesma civilização."),
    ("O QUE VARIA", "Vestígios, materiais, funções, períodos e relações com as descobertas anteriores."),
]
for idx, (label, body) in enumerate(axis_cards):
    y = 1.30 + idx * 1.75
    add_rect(slide, 0.55, y, 8.35, 1.40, PANEL, COPPER_DARK, radius=True)
    add_text(slide, f"0{idx + 1}", 0.80, y + 0.25, 0.60, 0.40, 15, COPPER, True, TITLE_FONT)
    add_text(slide, label, 1.52, y + 0.22, 2.35, 0.30, 11, SAND, True)
    add_text(slide, body, 1.52, y + 0.58, 6.95, 0.55, 14.2, STONE)


# 4 — Processo criativo
slide = new_slide()
add_slide_title(slide, "Processo criativo", "Um ciclo de geração, avaliação e acúmulo", 4)
steps = [
    "Conhecimento\nacumulado",
    "Direção\ndo grupo",
    "Prompt e\nreferências",
    "Geração\ncom IA",
    "Avaliação",
    "Aceitar, iterar\nou descartar",
    "Novo\nconhecimento",
]
positions = [(0.55, 1.50), (3.00, 1.50), (5.45, 1.50), (7.90, 1.50), (10.35, 1.50), (7.90, 3.38), (5.45, 3.38)]
for idx, (label, (x, y)) in enumerate(zip(steps, positions), 1):
    add_rect(slide, x, y, 2.15, 1.05, PANEL if idx != 7 else COPPER_DARK, COPPER_DARK, radius=True)
    add_text(slide, f"{idx}", x + 0.16, y + 0.14, 0.28, 0.28, 9, COPPER if idx != 7 else STONE, True)
    add_text(slide, label, x + 0.42, y + 0.22, 1.56, 0.56, 12.2, STONE, True, align=PP_ALIGN.CENTER, valign=MSO_ANCHOR.MIDDLE)
for a, b in zip(positions[:4], positions[1:5]):
    add_line(slide, a[0] + 2.15, a[1] + 0.52, b[0], b[1] + 0.52, COPPER, 1.3)
add_line(slide, 11.42, 2.55, 11.42, 3.38, COPPER, 1.3)
add_line(slide, 10.35, 3.90, 10.05, 3.90, COPPER, 1.3)
add_line(slide, 7.90, 3.90, 7.60, 3.90, COPPER, 1.3)
add_rect(slide, 0.60, 5.15, 5.65, 1.30, PANEL_2, PANEL_2, radius=True)
add_text(slide, "GRUPO", 0.86, 5.39, 0.90, 0.24, 10, COPPER, True)
add_text(slide, "Escolhe direções, constrói prompts, avalia e decide.", 1.75, 5.33, 4.15, 0.60, 13.4, STONE)
add_rect(slide, 6.55, 5.15, 5.90, 1.30, PANEL_2, PANEL_2, radius=True)
add_text(slide, "IA", 6.82, 5.39, 0.55, 0.24, 10, COPPER, True)
add_text(slide, "Produz possibilidades visuais para exploração e seleção.", 7.42, 5.33, 4.66, 0.60, 13.4, STONE)


# 5–8 — Descobertas
discovery_slides = [
    (
        5,
        "Descobertas 01–03",
        [
            (1, "Estela das Três Correntes", "Introduziu inscrições, pedra, cobre, reparo e o motivo das três linhas convergentes."),
            (2, "O Limiar Apagado", "Confirmou o motivo na arquitetura e revelou que uma das três partes foi apagada intencionalmente."),
            (3, "A Câmara do Terceiro Vazio", "Mostrou dois discos preservados e um terceiro receptáculo deliberadamente selado."),
        ],
    ),
    (
        6,
        "Descobertas 04–06",
        [
            (4, "O Disco Oculto", "Indicou que a terceira corrente não foi destruída: foi retirada e preservada em segredo."),
            (5, "A Mesa das Três Ofertas", "Levou a estrutura tripla a uma prática material e manteve o tratamento diferente da terceira parte."),
            (6, "O Fragmento das Três Rotas", "Sugeriu uma organização espacial e mostrou que a perda da terceira linha podia ser natural."),
        ],
    ),
    (
        7,
        "Descobertas 07–09",
        [
            (7, "A Confluência Vazia", "Reuniu trajetos, inscrições e um centro vazio, aproximando as ideias de rota e convergência."),
            (8, "O Tecido das Três Tramas", "Levou o padrão triplo da pedra para um objeto flexível, transportável e reparado."),
            (9, "O Cesto de Travessia", "Confirmou uma função prática para as tramas e introduziu o transporte de materiais diferentes."),
        ],
    ),
    (
        8,
        "Descobertas 10–12",
        [
            (10, "Os Recipientes do Cesto", "Revelou utensílios muito usados para manipular materiais, sem definir alimentação ou medição."),
            (11, "A Grade dos Quatro Espaços", "Mostrou que três divisórias podiam produzir quatro áreas e mudou a leitura do padrão triplo."),
            (12, "A Placa das Quatro Partes", "Conectou inscrições, cobre, três sulcos, quatro campos e transporte em um mesmo sistema material."),
        ],
    ),
]
for number, title, cards in discovery_slides:
    slide = new_slide()
    add_slide_title(slide, "Coleção", title, number)
    for card_index, (index, name, summary) in enumerate(cards):
        add_discovery_card(slide, index, name, summary, 0.45 + card_index * 4.28)


# 9 — Progressão
slide = new_slide()
add_slide_title(slide, "Progressão", "Como o conhecimento mudou ao longo da coleção", 9)
milestones = [
    (1, "SÍMBOLO", "01", "Três correntes, inscrições e um centro vazio."),
    (4, "OCULTAÇÃO", "02–04", "Apagamento, selamento e preservação secreta."),
    (6, "ESPAÇO", "05–07", "Práticas, possíveis rotas e convergências."),
    (9, "COTIDIANO", "08–10", "Tecido, transporte e utensílios muito usados."),
    (12, "REINTERPRETAÇÃO", "11–12", "Três elementos também podiam organizar quatro partes."),
]
add_line(slide, 1.15, 3.78, 12.05, 3.78, COPPER_DARK, 3)
for idx, (image_index, label, span, body) in enumerate(milestones):
    x = 0.46 + idx * 2.55
    add_image(slide, IMAGES[image_index], x, 1.33, 2.20, 1.76, darken=0.90, border=RGBColor(65, 77, 67))
    add_rect(slide, x + 0.87, 3.57, 0.45, 0.45, COPPER, COPPER, radius=True)
    add_text(slide, span, x, 4.18, 2.20, 0.25, 10, COPPER, True, align=PP_ALIGN.CENTER)
    add_text(slide, label, x, 4.54, 2.20, 0.30, 11, SAND, True, align=PP_ALIGN.CENTER)
    add_text(slide, body, x + 0.04, 5.04, 2.12, 1.03, 10.5, STONE, align=PP_ALIGN.CENTER)
add_text(slide, "Pontos de virada", 0.55, 6.46, 1.45, 0.25, 10, MUTED, True)
add_text(slide, "02 apagamento  ·  04 preservação  ·  09 uso cotidiano  ·  11 três ≠ três partes  ·  12 síntese", 2.10, 6.42, 10.35, 0.35, 11.2, COPPER)


# 10 — Descarte
slide = new_slide()
add_slide_title(slide, "Descarte comentado", "Uma imagem plausível ainda pode quebrar a continuidade", 10)
add_image(slide, DISCARD, 0.55, 1.34, 3.56, 4.42, darken=0.88, border=RUST)
add_text(slide, "DESCARTADA", 0.73, 1.51, 1.20, 0.25, 9, RUST, True)
add_text(slide, "Cozinha / preparo de alimentos", 0.55, 5.91, 3.56, 0.34, 13, STONE, True, TITLE_FONT, align=PP_ALIGN.CENTER)
add_text(slide, "→", 4.35, 3.08, 0.55, 0.55, 24, COPPER, True, align=PP_ALIGN.CENTER)
add_image(slide, IMAGES[10], 5.12, 1.34, 3.56, 4.42, darken=0.94, border=COPPER)
add_text(slide, "ACEITA", 5.30, 1.51, 0.80, 0.25, 9, COPPER, True)
add_text(slide, "Os Recipientes do Cesto", 5.12, 5.91, 3.56, 0.34, 13, STONE, True, TITLE_FONT, align=PP_ALIGN.CENTER)
add_rect(slide, 9.02, 1.34, 3.78, 4.92, PANEL, COPPER_DARK, radius=True)
add_text(slide, "POR QUE NÃO ENTROU?", 9.30, 1.70, 3.15, 0.30, 10.5, RUST, True)
reasons = [
    "Tratava os resíduos como alimentos sem comprovação.",
    "Pressupunha uma cozinha próxima ao cesto.",
    "Criava uma relação funcional ainda não sustentada.",
]
for idx, reason in enumerate(reasons, 1):
    y = 2.22 + (idx - 1) * 0.82
    add_text(slide, f"0{idx}", 9.30, y, 0.38, 0.25, 9, COPPER, True)
    add_text(slide, reason, 9.75, y - 0.03, 2.65, 0.55, 11.7, STONE)
add_line(slide, 9.30, 4.80, 12.42, 4.80, COPPER_DARK, 1)
add_text(slide, "APRENDIZADO", 9.30, 5.04, 1.20, 0.25, 9, SAND, True)
add_text(slide, "Continuidade também limita quantas inferências entram de uma vez.", 9.30, 5.39, 3.03, 0.64, 12.2, COPPER, True)


# 11 — Ferramentas e IA
slide = new_slide()
add_slide_title(slide, "Ferramentas e IA", "A geração fazia parte de um processo de decisão", 11)
add_image(slide, IMAGES[1], 0.55, 1.35, 2.25, 2.80, darken=0.80, border=RGBColor(65, 77, 67))
add_image(slide, IMAGES[9], 2.98, 1.35, 2.25, 2.80, darken=0.80, border=RGBColor(65, 77, 67))
add_rect(slide, 0.55, 4.42, 4.68, 1.57, PANEL, COPPER_DARK, radius=True)
add_text(slide, "FERRAMENTAS DOCUMENTADAS", 0.82, 4.70, 3.85, 0.25, 9.5, COPPER, True)
add_text(slide, "ChatGPT · descobertas 01–08\nCodex + gerador integrado · 09–10", 0.82, 5.08, 3.95, 0.62, 13, STONE, True, line_spacing=0.95)
add_rect(slide, 5.70, 1.35, 6.98, 4.64, PANEL, PANEL, radius=True)
add_text(slide, "GRUPO", 6.05, 1.73, 1.05, 0.26, 10, SAND, True)
add_text(slide, "direção  ·  contexto  ·  prompt", 7.15, 1.69, 4.70, 0.33, 14, STONE, True)
add_line(slide, 6.05, 2.30, 12.20, 2.30, COPPER_DARK, 1)
add_text(slide, "IA", 6.05, 2.67, 1.05, 0.26, 10, SAND, True)
add_text(slide, "possibilidades visuais  ·  variações", 7.15, 2.63, 4.70, 0.33, 14, STONE, True)
add_line(slide, 6.05, 3.24, 12.20, 3.24, COPPER_DARK, 1)
add_text(slide, "GRUPO", 6.05, 3.61, 1.05, 0.26, 10, SAND, True)
add_text(slide, "avaliação  ·  iteração  ·  descarte", 7.15, 3.57, 4.70, 0.33, 14, STONE, True)
add_line(slide, 6.05, 4.18, 12.20, 4.18, COPPER_DARK, 1)
add_text(slide, "COLEÇÃO", 6.05, 4.55, 1.05, 0.26, 10, COPPER, True)
add_text(slide, "somente resultados aceitos viram contexto", 7.15, 4.51, 4.90, 0.48, 14, COPPER, True)
add_text(slide, "Imagens anteriores foram usadas como referência visual; seed e parâmetros nem sempre estavam disponíveis.", 6.05, 5.30, 5.96, 0.48, 10.6, MUTED)


# 12 — Skills
slide = new_slide()
add_slide_title(slide, "Skills utilizadas", "Três intervenções que mudaram a proposta", 12)
skill_cards = [
    (1, "abrir-o-leque", "Do museu de objetos independentes para uma sequência de descobertas."),
    (2, "afiar-o-eixo", "Definição da restrição, da repetição, da variação e do pertencimento."),
    (10, "derrubar-a-ideia", "Continuidade, ordem, descarte e participação ativa do grupo."),
]
for idx, (image_index, name, body) in enumerate(skill_cards):
    x = 0.50 + idx * 4.28
    add_rect(slide, x, 1.35, 4.02, 5.15, PANEL, COPPER_DARK, radius=True)
    add_image(slide, IMAGES[image_index], x + 0.10, 1.45, 3.82, 2.28, darken=0.72, border=RGBColor(64, 76, 67))
    add_text(slide, f"`{name}`", x + 0.23, 4.06, 3.55, 0.42, 15.5, COPPER, True, BODY_FONT, align=PP_ALIGN.CENTER)
    add_text(slide, body, x + 0.30, 4.72, 3.40, 1.05, 13, STONE, align=PP_ALIGN.CENTER)
add_text(slide, "A skill `escutar-a-reuniao` permanece registrada como não utilizada.", 0.55, 6.68, 12.15, 0.27, 9.5, MUTED, align=PP_ALIGN.CENTER)


# 13 — Aprendizados
slide = new_slide()
add_slide_title(slide, "Aprendizados", "O que o processo tornou visível", 13)
lessons = [
    (1, "Coerência não é só estética", "Repetir pedra, cobre e símbolos não basta: cada peça precisa alterar o conhecimento."),
    (8, "Evidência não é interpretação", "As descrições distinguem o que foi observado das hipóteses ainda abertas."),
    (10, "Descartar também constrói", "Inserir etapas intermediárias evitou saltos e tornou a progressão mais defensável."),
]
for idx, (image_index, title, body) in enumerate(lessons):
    y = 1.32 + idx * 1.78
    add_image(slide, IMAGES[image_index], 0.58, y, 2.08, 1.48, darken=0.78, border=RGBColor(64, 76, 67))
    add_rect(slide, 2.90, y, 9.83, 1.48, PANEL, COPPER_DARK, radius=True)
    add_text(slide, f"0{idx + 1}", 3.20, y + 0.24, 0.48, 0.30, 12, COPPER, True, TITLE_FONT)
    add_text(slide, title, 3.83, y + 0.20, 3.65, 0.34, 16, STONE, True, TITLE_FONT)
    add_text(slide, body, 3.83, y + 0.68, 8.24, 0.48, 12.5, MUTED)


# 14 — Encerramento
slide = new_slide()
add_image(slide, IMAGES[1], 0, 0, 6.666, 7.5, darken=0.56, center=(0.5, 0.47))
add_image(slide, IMAGES[12], 6.666, 0, 6.667, 7.5, darken=0.56, center=(0.5, 0.48))
add_rect(slide, 2.02, 1.52, 9.30, 4.35, BG, BG, radius=True)
add_text(slide, "DA PRIMEIRA PISTA À ÚLTIMA CONEXÃO", 2.55, 1.95, 8.25, 0.30, 10, COPPER, True, align=PP_ALIGN.CENTER)
add_text(
    slide,
    "A civilização não foi definida antes das imagens.",
    2.62,
    2.53,
    8.10,
    0.62,
    24,
    STONE,
    True,
    TITLE_FONT,
    align=PP_ALIGN.CENTER,
)
add_text(slide, "Ela emergiu das relações construídas entre elas.", 2.72, 3.25, 7.90, 0.46, 18, SAND, align=PP_ALIGN.CENTER)
add_line(slide, 4.22, 4.03, 9.12, 4.03, COPPER_DARK, 1.2)
add_text(
    slide,
    "As três correntes registravam lugares, materiais, grupos, ações — ou uma forma de organizar relações?",
    3.02,
    4.39,
    7.30,
    0.92,
    14.2,
    COPPER,
    True,
    align=PP_ALIGN.CENTER,
)
add_text(slide, "ECOS DE UMA CIVILIZAÇÃO", 4.30, 6.72, 4.75, 0.30, 9, STONE, True, align=PP_ALIGN.CENTER)


prs.core_properties.title = "Ecos de uma Civilização"
prs.core_properties.subject = "Apresentação do projeto de Criatividade Computacional"
prs.core_properties.author = "Equipe Ecos de uma Civilização"
prs.core_properties.keywords = "arqueologia fictícia, criatividade computacional, IA, coleção"
prs.core_properties.comments = "Gerado exclusivamente a partir do conteúdo do repositório do projeto."

prs.save(OUT)
print(f"Apresentação criada: {OUT}")
print(f"Slides: {len(prs.slides)}")

